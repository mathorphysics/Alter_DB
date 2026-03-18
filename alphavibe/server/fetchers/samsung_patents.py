"""
Samsung Semiconductor Patent Filing Trend
Parses manually exported KIPRIS 분류통계 (classification statistics) Excel file.

Export steps:
  1. Go to https://www.kipris.or.kr/khome/search/searchResult.do?tab=patent
  2. Advanced search: 출원인 = 삼성전자, 분류번호 = H01L
  3. Export → 분류통계 Excel

File format (one row = one rank entry per column):
  출원년도       공개년도   등록년도   IPC          CPC          출원인              발명자
  2015(191)    2017(118)  2022(125)  H01L(369)   H01L(500)   삼성전자주식회사(958) 박영우(39)
  ...

We extract the 출원년도 (application year) column → annual patent count time series.

Cached to: server/data/samsung_patents.json
"""

import json
import logging
import re
import zipfile
import io
from pathlib import Path
from typing import Dict, List, Optional

logger = logging.getLogger(__name__)

DATA_DIR   = Path(__file__).parent.parent / "data"
CACHE_PATH = DATA_DIR / "samsung_patents.json"
XLSX_GLOB  = "KIPRIS*.xlsx"


def _find_latest_xlsx() -> Optional[Path]:
    files = sorted(DATA_DIR.glob(XLSX_GLOB))
    return files[-1] if files else None


def _load_workbook_patched(path: Path):
    """Load xlsx while patching invalid RGB color values that KIPRIS generates."""
    import openpyxl

    with zipfile.ZipFile(path, "r") as z:
        styles_raw = z.read("xl/styles.xml").decode("utf-8")

    fixed = re.sub(
        r'rgb="([^"]+)"',
        lambda m: 'rgb="FF000000"'
        if not re.match(r"^[0-9A-Fa-f]{6,8}$", m.group(1))
        else m.group(0),
        styles_raw,
    )

    buf = io.BytesIO()
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/styles.xml":
                data = fixed.encode("utf-8")
            zout.writestr(item, data)

    buf.seek(0)
    return openpyxl.load_workbook(buf)


def _parse_year_count(cell: str) -> Optional[tuple[int, int]]:
    """'2015(191)' → (2015, 191)"""
    if not cell:
        return None
    m = re.match(r"(\d{4})\((\d+)\)", str(cell).strip())
    if not m:
        return None
    return int(m.group(1)), int(m.group(2))


def parse_kipris_xlsx(path: Path) -> List[Dict]:
    """
    Parse KIPRIS 분류통계 Excel → annual patent application counts.
    Returns list sorted by year: [{"period": "2015", "label": "2015", "count": 191}, ...]
    """
    wb = _load_workbook_patched(path)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Empty worksheet")

    # Find the 출원년도 column index from header row
    header = [str(c).strip() if c else "" for c in rows[0]]
    try:
        col_idx = header.index("출원년도")
    except ValueError:
        # Fallback: first column
        col_idx = 0
        logger.warning("출원년도 column not found, using column 0")

    counts: Dict[int, int] = {}
    for row in rows[1:]:
        if col_idx >= len(row):
            continue
        parsed = _parse_year_count(row[col_idx])
        if parsed:
            year, count = parsed
            counts[year] = count

    result = sorted(
        [{"period": str(y), "label": str(y), "count": c} for y, c in counts.items()],
        key=lambda r: r["period"],
    )

    logger.info("Parsed %d years from %s", len(result), path.name)
    return result


def load_patents() -> Optional[Dict]:
    return json.loads(CACHE_PATH.read_text()) if CACHE_PATH.exists() else None


def get_patents() -> Dict:
    cached = load_patents()
    if cached:
        return cached
    return refresh_from_xlsx()


def refresh_from_xlsx() -> Dict:
    xlsx = _find_latest_xlsx()
    if not xlsx:
        raise FileNotFoundError(
            f"No KIPRIS xlsx found in {DATA_DIR}. "
            "Export 분류통계 from https://www.kipris.or.kr and place in server/data/"
        )
    timeseries = parse_kipris_xlsx(xlsx)
    output = {"timeseries": timeseries, "source_file": xlsx.name}

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    CACHE_PATH.write_text(json.dumps(output, ensure_ascii=False, indent=2))
    logger.info("Cached %d years → %s", len(timeseries), CACHE_PATH)
    return output


if __name__ == "__main__":
    import sys
    logging.basicConfig(level=logging.INFO,
                        format="%(asctime)s  %(levelname)-7s  %(message)s")
    try:
        data = refresh_from_xlsx()
        ts = data["timeseries"]
        print(f"\n{len(ts)} years parsed from {data['source_file']}\n")
        print(f"{'Year':<8}  {'Count':>8}")
        print("-" * 20)
        for row in ts:
            print(f"{row['period']:<8}  {row['count']:>8}")
    except Exception as e:
        logger.error("Failed: %s", e)
        sys.exit(1)
